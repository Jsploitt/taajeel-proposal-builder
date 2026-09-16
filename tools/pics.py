"""Map picture shapes -> their media parts, for the slides that carry client assets."""
import re, zipfile, sys, json
from pptx import Presentation

SRC = "template/taajeel_master.pptx"
prs = Presentation(SRC)
z = zipfile.ZipFile(SRC)

def walk(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == 6:
            yield from walk(sh.shapes)

for n in [int(a) for a in sys.argv[1:]] or [1, 12, 14, 27]:
    slide = prs.slides[n - 1]
    print(f"=== slide {n}")
    for sh in walk(slide.shapes):
        if sh.shape_type != 13:   # PICTURE
            continue
        try:
            part = sh.image.blob
            fn = sh.image.filename or "?"
            sz = len(part)
            px = sh.image.size
        except Exception as e:
            print(f"   id={sh.shape_id} {sh.name!r}  <no image> {e}")
            continue
        print(f"   id={sh.shape_id:<5} {sh.name!r:<22} {sh.image.content_type:<12} "
              f"{px[0]}x{px[1]}px {sz//1024}KB  pos=({round(sh.left/914400,2)},{round(sh.top/914400,2)}) "
              f"size=({round(sh.width/914400,2)}x{round(sh.height/914400,2)}in)  sha={sh.image.sha1[:8]}")
