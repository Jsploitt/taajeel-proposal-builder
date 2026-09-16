"""Run-level dump for a given slide/shape. Exposes where a single visual line
is split across runs -- the thing that makes find-and-replace fail."""
import sys
from pptx import Presentation
from pptx.util import Pt

SRC = "proposal/Proposal to HAVENSTONE CONSULTING W.L.L foreign company set up.pptx"
slide_n = int(sys.argv[1]); shape_id = int(sys.argv[2]) if len(sys.argv) > 2 else None

prs = Presentation(SRC)
slide = prs.slides[slide_n - 1]

def walk(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == 6:
            yield from walk(sh.shapes)

for sh in walk(slide.shapes):
    if shape_id and sh.shape_id != shape_id:
        continue
    if not sh.has_text_frame:
        continue
    print(f"--- shape id={sh.shape_id} name={sh.name!r}")
    for pi, p in enumerate(sh.text_frame.paragraphs):
        algn = p.alignment
        print(f"  p{pi:02d} algn={algn} lvl={p.level} runs={len(p.runs)}")
        for ri, r in enumerate(p.runs):
            f = r.font
            col = None
            try:
                col = f.color.rgb
            except Exception:
                col = f.color.type
            print(f"      r{ri} sz={f.size.pt if f.size else None} b={f.bold} "
                  f"i={f.italic} font={f.name!r} color={col} :: {r.text!r}")
